"""Import the registrar's SIS equivalency export (KAN-42).

The export is one row per approved (external course -> Lipscomb course) pairing that has been
used at least 3 times, with these columns (names as they come out of the SIS query):

    Id, College Course.College.Code, College Course.College.Name, College Course.Code,
    College Course.Name, College Course.Credits, College Course.Minimum Grade Required,
    Course.Code, Course.Description, Course.Credit Hours, TimesTransferred

Two halves:
- parse_rows() / build_plan() are pure Python: rows in, an ImportPlan out. Unit tested directly.
- load_sis_equivalencies() writes a plan into the database. Idempotent per external course.

How rows become equivalencies (see specs/03-design-equivalencies.md, "SIS import"):
- Every pairing in the export was approved by the registrar's office, so the most-used Lipscomb
  course for an external course is stored as `confirmed` with decided_by = REGISTRAR_SIS.
- Some external courses were approved as more than one Lipscomb course. Only one row per external
  course may be `confirmed`, so the others are stored as `proposed` (source sis_import) with
  confidence = their share of transfers, for the registrar to sort out.
- Lipscomb codes with X placeholders (EN2XXT, PH1XXX) and TRACC are elective credit: kind
  `elective`, no course_id, the code kept in notes.

The real export is not committed: the repo is public. Keep it in data/private/ (gitignored).
"""
from __future__ import annotations

import csv
import re
from collections import defaultdict
from dataclasses import dataclass, field
from pathlib import Path

REGISTRAR_SIS = "Registrar's office (SIS)"
SIS_CATALOG_YEAR = "SIS"  # the export has no catalog year

COLUMNS = {
    "school_code": "College Course.College.Code",
    "school": "College Course.College.Name",
    "ext_code": "College Course.Code",
    "ext_title": "College Course.Name",
    "ext_credits": "College Course.Credits",
    "min_grade": "College Course.Minimum Grade Required",
    "lip_code": "Course.Code",
    "lip_title": "Course.Description",
    "lip_credits": "Course.Credit Hours",
    "times": "TimesTransferred",
}

_CODE = re.compile(r"^([A-Z]+)[\s\-]*(.*)$")


@dataclass(frozen=True)
class SisRow:
    school: str
    school_code: str
    ext_code: str
    ext_title: str
    ext_credits: float | None
    min_grade: str
    lip_code: str
    lip_title: str
    lip_credits: int | None
    times: int


@dataclass
class Mapping:
    """One Lipscomb course (or elective code) an external course was approved as."""
    lip_code: str
    lip_subject: str
    lip_number: str
    lip_title: str
    lip_credits: int | None
    elective: bool
    times: int = 0
    min_grades: set[str] = field(default_factory=set)


@dataclass
class PlannedExternal:
    school: str
    subject: str
    number: str
    title: str
    credits: float | None
    mappings: list[Mapping]  # most-transferred first; [0] is the one stored as confirmed

    @property
    def total_times(self) -> int:
        return sum(m.times for m in self.mappings)


@dataclass
class ImportPlan:
    schools: dict[str, str]  # name -> kind
    externals: list[PlannedExternal]
    lipscomb_courses: dict[tuple[str, str], tuple[str, int | None]]  # (subj, num) -> (title, credits)

    def summary(self) -> dict[str, int]:
        return {
            "rows_schools": len(self.schools),
            "rows_external_courses": len(self.externals),
            "rows_lipscomb_courses": len(self.lipscomb_courses),
            "confirmed": len(self.externals),
            "proposed_alternates": sum(len(e.mappings) - 1 for e in self.externals),
            "multi_mapped_external_courses": sum(len(e.mappings) > 1 for e in self.externals),
            "electives": sum(m.elective for e in self.externals for m in e.mappings),
        }


# --- pure helpers ----------------------------------------------------------------------------

def split_code(code: object) -> tuple[str, str]:
    """'ENGL 1010' / 'ENGL1010' / 'AR-2201-0399-03' -> (subject, number).

    Codes with no leading letters (e.g. '3750100') come back as ('', code); codes that are all
    letters (e.g. 'TRACC') come back as (code, '').
    """
    s = re.sub(r"\s+", " ", str(code).strip().upper())
    m = _CODE.match(s)
    if not m:
        return "", s
    return m.group(1), m.group(2).strip()


def is_elective_code(code: str) -> bool:
    """Lipscomb placeholder codes: 'X' in the number (EN2XXT, PH1XXX, RTC12XT) or bare 'TRACC'."""
    subject, number = split_code(code)
    return not number or "X" in number


def school_kind(name: str) -> str:
    return "community_college" if re.search(r"community college|technical|tech\b", name, re.I) else "university"


def _num(v) -> float | None:
    if v is None or str(v).strip() == "":
        return None
    f = float(v)
    return int(f) if f.is_integer() else round(f, 2)


def parse_rows(records: list[dict]) -> list[SisRow]:
    """Header-keyed dicts (from xlsx or csv) -> SisRow. Raises ValueError on a missing column."""
    if not records:
        return []
    missing = [col for col in COLUMNS.values() if col not in records[0]]
    if missing:
        raise ValueError(f"SIS export is missing columns: {', '.join(missing)}")
    rows = []
    for r in records:
        if not str(r[COLUMNS["school"]] or "").strip() or not str(r[COLUMNS["lip_code"]] or "").strip():
            continue
        lip_credits = _num(r[COLUMNS["lip_credits"]])
        rows.append(SisRow(
            school=str(r[COLUMNS["school"]]).strip(),
            school_code=str(r[COLUMNS["school_code"]] or "").strip(),
            ext_code=str(r[COLUMNS["ext_code"]]).strip(),
            ext_title=str(r[COLUMNS["ext_title"]] or "").strip(),
            ext_credits=_num(r[COLUMNS["ext_credits"]]),
            min_grade=str(r[COLUMNS["min_grade"]] or "").strip(),
            lip_code=re.sub(r"\s+", "", str(r[COLUMNS["lip_code"]]).upper()),
            lip_title=str(r[COLUMNS["lip_title"]] or "").strip(),
            lip_credits=int(lip_credits) if lip_credits is not None else None,
            times=int(_num(r[COLUMNS["times"]]) or 0),
        ))
    return rows


def build_plan(rows: list[SisRow]) -> ImportPlan:
    schools: dict[str, str] = {}
    lipscomb: dict[tuple[str, str], tuple[str, int | None]] = {}
    # (school, subject, number) -> lip_code -> Mapping; plus the best title/credits seen
    grouped: dict[tuple[str, str, str], dict[str, Mapping]] = defaultdict(dict)
    ext_best: dict[tuple[str, str, str], tuple[int, str, float | None]] = {}

    for r in rows:
        schools.setdefault(r.school, school_kind(r.school))
        subject, number = split_code(r.ext_code)
        key = (r.school, subject, number)
        if key not in ext_best or r.times > ext_best[key][0]:
            ext_best[key] = (r.times, r.ext_title or r.ext_code, r.ext_credits)

        elective = is_elective_code(r.lip_code)
        lsubj, lnum = split_code(r.lip_code)
        m = grouped[key].get(r.lip_code)
        if m is None:
            m = grouped[key][r.lip_code] = Mapping(r.lip_code, lsubj, lnum, r.lip_title, r.lip_credits, elective)
        m.times += r.times
        if r.min_grade:
            m.min_grades.add(r.min_grade)
        if not elective:
            lipscomb.setdefault((lsubj, lnum), (r.lip_title or r.lip_code, r.lip_credits))

    externals = []
    for key, by_code in grouped.items():
        school, subject, number = key
        _, title, credits = ext_best[key]
        # Most transfers first; ties go to a real course over an elective, then to the shorter code.
        mappings = sorted(by_code.values(), key=lambda m: (-m.times, m.elective, len(m.lip_code), m.lip_code))
        externals.append(PlannedExternal(school, subject, number, title, credits, mappings))
    externals.sort(key=lambda e: (e.school, e.subject, e.number))
    return ImportPlan(schools, externals, lipscomb)


def read_export(path: str | Path) -> list[dict]:
    """Read the export (.xlsx or .csv) into header-keyed dicts."""
    p = Path(path)
    if p.suffix.lower() == ".csv":
        with open(p, newline="", encoding="utf-8-sig") as f:
            return list(csv.DictReader(f))
    from openpyxl import load_workbook  # dev/CLI only; not needed by the deployed API

    wb = load_workbook(p, read_only=True, data_only=True)
    try:
        for ws in wb.worksheets:
            it = ws.iter_rows(values_only=True)
            header = next(it, None)
            if header and COLUMNS["lip_code"] in header:
                names = [str(h) if h is not None else "" for h in header]
                return [dict(zip(names, row)) for row in it]
    finally:
        wb.close()
    raise ValueError(f"no sheet in {p.name} has a {COLUMNS['lip_code']!r} column")


# --- database --------------------------------------------------------------------------------

def _notes(m: Mapping, total: int, primary: bool) -> str:
    parts = []
    if m.elective:
        parts.append(f"Lipscomb elective code {m.lip_code}")
    grades = ", ".join(sorted(m.min_grades)) or "not given"
    parts.append(f"approved in SIS, used {m.times} time{'s' if m.times != 1 else ''}, min grade {grades}")
    if not primary:
        parts.append(f"also approved for this course ({m.times} of {total} transfers); registrar to pick")
    text = "; ".join(parts)
    return text[0].upper() + text[1:]


def load_sis_equivalencies(path: str | Path) -> dict[str, int]:
    """Load the SIS export into the database. Needs an app context.

    Idempotent per external course: an external course that already has any sis_import row is
    skipped, so re-running the same file adds nothing. (Updating changed rows is not supported
    yet; that's a follow-up once we know how often the registrar re-exports.)
    """
    from .models import Course, Equivalency, ExternalCourse, School, db

    plan = build_plan(parse_rows(read_export(path)))
    counts = {"schools": 0, "external_courses": 0, "lipscomb_courses": 0, "confirmed": 0,
              "proposed_alternates": 0, "skipped_existing": 0}

    schools = {s.name: s for s in School.query.all()}
    for name, kind in plan.schools.items():
        if name not in schools:
            schools[name] = School(name=name, kind=kind)
            db.session.add(schools[name])
            counts["schools"] += 1

    courses = {(c.subject, c.number): c for c in Course.query.all()}
    for (subj, num), (title, credits) in plan.lipscomb_courses.items():
        if (subj, num) not in courses:
            courses[(subj, num)] = Course(subject=subj, number=num, title=title[:120],
                                             credits=credits if credits is not None else 0)
            db.session.add(courses[(subj, num)])
            counts["lipscomb_courses"] += 1
    db.session.flush()

    imported = {eid for (eid,) in db.session.query(Equivalency.external_course_id)
                .filter(Equivalency.source == "sis_import").distinct()}
    externals = {(e.school_id, e.subject, e.number): e
                 for e in ExternalCourse.query.filter_by(catalog_year=SIS_CATALOG_YEAR)}

    for pe in plan.externals:
        school = schools[pe.school]
        ext = externals.get((school.id, pe.subject, pe.number))
        if ext is None:
            ext = ExternalCourse(school=school, subject=pe.subject, number=pe.number, title=pe.title[:160],
                                 credits=pe.credits, catalog_year=SIS_CATALOG_YEAR)
            db.session.add(ext)
            counts["external_courses"] += 1
        elif ext.id in imported:
            counts["skipped_existing"] += 1
            continue

        for i, m in enumerate(pe.mappings):
            primary = i == 0
            db.session.add(Equivalency(
                external_course=ext,
                course=None if m.elective else courses[(m.lip_subject, m.lip_number)],
                kind="elective" if m.elective else "direct",
                status="confirmed" if primary else "proposed",
                source="sis_import",
                confidence=None if primary else round(m.times / pe.total_times, 2),
                decided_by=REGISTRAR_SIS if primary else None,
                notes=_notes(m, pe.total_times, primary),
            ))
            counts["confirmed" if primary else "proposed_alternates"] += 1

    db.session.commit()
    return counts
